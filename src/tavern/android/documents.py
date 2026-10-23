"""Office parsing in a disposable Android service, never in the API process."""

import asyncio
import io
import json

_conversion_lock = asyncio.Lock()


class ConversionLimit(ValueError):
    """The only converter error whose text reaches the caller unchanged."""


def convert_office(blob: bytes, extension: str, max_chars: int) -> str:
    """Called only after archive validation and inside the :documents process."""
    if extension == ".docx":
        import mammoth
        from markdownify import markdownify

        # No document relationship can read another local file or fetch an external image.
        converted = mammoth.convert_to_html(
            io.BytesIO(blob),
            external_file_access=False,
            convert_image=lambda _: [],
        )
        return markdownify(converted.value, heading_style="ATX", bullets="*").strip()

    from openpyxl import load_workbook

    result = io.StringIO()
    chars = 0

    def write(text: str) -> None:
        nonlocal chars
        chars += len(text)
        if chars > max_chars:
            raise ConversionLimit("转换后 Markdown 超过 800 万字符，请先拆分文件。")
        result.write(text)

    def cell(value) -> str:
        if value is None:
            return ""
        return (
            str(value)
            .replace("\\", "\\\\")
            .replace("|", "\\|")
            .replace("\r\n", "\n")
            .replace("\r", "\n")
            .replace("\n", "<br>")
        )

    workbook = load_workbook(io.BytesIO(blob), read_only=True, data_only=True, keep_links=False)
    try:
        for sheet in workbook.worksheets:
            write(f"## {sheet.title}\n\n")
            # Producers can understate worksheet dimensions. A bounded first scan finds
            # the real table width so no later cells disappear behind a short Markdown header.
            sheet.reset_dimensions()
            width = 0
            for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                width = max(width, len(row))
                if row_number * width * 3 > max_chars:
                    raise ConversionLimit("转换后 Markdown 超过 800 万字符，请先拆分文件。")
            if not width:
                continue
            for row_number, row in enumerate(sheet.iter_rows(values_only=True, max_col=width)):
                write("| " + " | ".join(cell(value) for value in row) + " |\n")
                if row_number == 0:
                    write("| " + " | ".join("---" for _ in range(width)) + " |\n")
            write("\n")
    finally:
        workbook.close()
    return result.getvalue().strip()


async def read_document(path: str, offset: int, limit: int, timeout: float) -> dict:
    from java import jclass

    # The service is deliberately single-use. Serializing conversions prevents cancellation
    # of one model tool call from terminating a different call's parser process.
    async with _conversion_lock:
        context = jclass("com.chaquo.python.Python").getPlatform().getApplication()
        client = jclass("com.terrarium.tavern.DocumentClient")(
            context, json.dumps({"path": path, "offset": offset, "limit": limit})
        )
        operation = asyncio.create_task(asyncio.to_thread(client.result, int(timeout * 1000)))
        try:
            response = await asyncio.shield(operation)
            result = json.loads(str(response))
            if "error" in result:
                raise ValueError(result["error"])
            return result
        finally:
            # close kills the separate process and completes the native future, including
            # cancellation during service binding. Drain its waiting thread before reuse.
            await asyncio.to_thread(client.close)
            await operation
