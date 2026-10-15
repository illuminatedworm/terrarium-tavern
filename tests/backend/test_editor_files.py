"""External references are bounded, resumable pages and never draft read evidence."""
import asyncio
import io
import json
import sys
import zipfile
from types import SimpleNamespace

import pytest
from openpyxl import Workbook
from pydantic import ValidationError

from tavern.character import editor_files
from tavern.character.card_editor import CardEditorService
from tavern.character.editor_files import convert_document, read_document, read_text_file
from tavern.character.editor_tools import DocumentFile, EditorTools, TextFile
from tavern.core.model_view import ModelView
from tavern.storage.store import Store
from tests.backend.test_card_authoring import run
from tests.backend.test_card_editor import ToolLLM, import_fixture
from tests.backend.test_card_editor_alignment import Editor
from tests.backend.test_half_agent_alignment import blocks


def make_docx(path):
    """A real, minimal OOXML document; no extra authoring dependency just for fixtures."""
    ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    files = {
        "[Content_Types].xml": '''<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
          <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
          <Default Extension="xml" ContentType="application/xml"/>
          <Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
        </Types>''',
        "_rels/.rels": '''<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
          <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
        </Relationships>''',
        "word/_rels/document.xml.rels": '''<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
          <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>
          <Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/numbering" Target="numbering.xml"/>
        </Relationships>''',
        "word/styles.xml": f'''<w:styles xmlns:w="{ns}"><w:style w:type="paragraph" w:styleId="Heading1">
          <w:name w:val="heading 1"/></w:style></w:styles>''',
        "word/numbering.xml": f'''<w:numbering xmlns:w="{ns}"><w:abstractNum w:abstractNumId="0">
          <w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="bullet"/><w:lvlText w:val="•"/></w:lvl>
          </w:abstractNum><w:num w:numId="1"><w:abstractNumId w:val="0"/></w:num></w:numbering>''',
        "word/document.xml": f'''<w:document xmlns:w="{ns}"><w:body>
          <w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>世界设定</w:t></w:r></w:p>
          <w:p><w:r><w:t>海港人物🌊</w:t></w:r></w:p>
          <w:p><w:pPr><w:numPr><w:ilvl w:val="0"/><w:numId w:val="1"/></w:numPr></w:pPr><w:r><w:t>航海经历</w:t></w:r></w:p>
          <w:tbl><w:tr><w:tc><w:p><w:r><w:t>姓名</w:t></w:r></w:p></w:tc>
          <w:tc><w:p><w:r><w:t>职业</w:t></w:r></w:p></w:tc></w:tr>
          <w:tr><w:tc><w:p><w:r><w:t>林澈</w:t></w:r></w:p></w:tc>
          <w:tc><w:p><w:r><w:t>航海士</w:t></w:r></w:p></w:tc></w:tr></w:tbl>
          <w:p><w:r><w:t>文档完整尾部</w:t></w:r></w:p></w:body></w:document>''',
    }
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, value in files.items():
            archive.writestr(name, value.encode("utf-8"))
    return path


def make_xlsx(path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "人物"
    sheet.append(["姓名", "说明"])
    sheet.append(["林澈", "海港人物"])
    sheet.append(["有缓存公式", "=1+2"])
    sheet.append(["无缓存公式", "=2+3"])
    sheet.append(["零值", 0])
    workbook.create_sheet("空表")
    workbook.create_sheet("世界").append(["文档完整尾部"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    with zipfile.ZipFile(buffer) as source, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            data = source.read(info)
            if info.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(b"<f>1+2</f><v></v>", b"<f>1+2</f><v>3</v>")
            target.writestr(info, data)
    return path


@pytest.mark.parametrize("encoding", ["utf-8", "utf-8-sig", "utf-16", "utf-32", "gb18030"])
def test_text_pages_preserve_encoding_newlines_and_eof(tmp_path, encoding):
    path = tmp_path / "资料.csv"
    original = '姓名,备注\r\n林澈,"海港"\r\n\r\n空行之后\r最后一行\n末行无换行'
    path.write_bytes(original.encode(encoding))
    offset, pages = 1, []
    while True:
        result = read_text_file(str(path), offset=offset, limit=2,
                                encoding=encoding if encoding == "gb18030" else None)
        pages.append(result["text"])
        assert result["returned_lines"] <= 2
        if not result["has_more"]:
            assert result["next_offset"] is None
            break
        assert result["next_offset"] == offset + result["returned_lines"]
        offset = result["next_offset"]
    assert "".join(pages) == original
    assert path.read_bytes() == original.encode(encoding)
    empty = read_text_file(str(path), offset=999, encoding=encoding)
    assert empty["text"] == "" and not empty["has_more"] and empty["next_offset"] is None


def test_text_default_and_character_limits_are_complete_lines(tmp_path):
    path = tmp_path / "大文件.txt"
    path.write_text("一行\n" * 201, encoding="utf-8")
    first = read_text_file(str(path))
    assert first["returned_lines"] == 200 and first["next_offset"] == 201
    path.write_bytes(("甲" * 9999 + "\r\n" + "乙" * 9999 + "\n" + "尾部").encode())
    first = read_text_file(str(path))
    second = read_text_file(str(path), offset=first["next_offset"])
    assert first["returned_lines"] == 1
    assert first["text"] + second["text"] == path.read_bytes().decode()
    path.write_text("甲" * 20_000, encoding="utf-8")
    assert read_text_file(str(path))["text"] == "甲" * 20_000
    path.write_text("开头\n" + "乙" * 20_001, encoding="utf-8")
    assert read_text_file(str(path))["next_offset"] == 2
    with pytest.raises(ValueError, match="第 2 行超过"):
        read_text_file(str(path), offset=2)


def test_text_empty_binary_encoding_and_source_limits(tmp_path):
    path = tmp_path / "资料.txt"
    path.write_bytes(b"")
    assert read_text_file(str(path))["returned_lines"] == 0
    path.write_bytes(b"hello\0world")
    with pytest.raises(ValueError, match="二进制"):
        read_text_file(str(path))
    path.write_bytes("中文".encode("gb18030"))
    with pytest.raises(ValueError, match="解码失败"):
        read_text_file(str(path))
    for encoding in ("not-an-encoding", "base64_codec"):
        with pytest.raises(ValueError, match="encoding"):
            read_text_file(str(path), encoding=encoding)
    with path.open("wb") as stream:
        stream.truncate(editor_files.MAX_FILE_BYTES + 1)
    with pytest.raises(ValueError, match="32 MiB"):
        read_text_file(str(path))
    # Enforce actual consumed bytes as well as the stat check, without a racing fixture.
    reader = editor_files._BoundedReader(io.BytesIO(b"ab"))
    reader.consumed = editor_files.MAX_FILE_BYTES - 1
    with pytest.raises(ValueError, match="32 MiB"):
        reader.readinto(bytearray(2))


@pytest.mark.parametrize("value", ["relative.txt", "https://example.com/a.docx", "file:///tmp/a.txt",
                                 r"\\server\share\a.txt", r"\\.\NUL", "/dev/null"])
def test_nonlocal_or_device_paths_rejected(value):
    with pytest.raises(ValueError):
        read_text_file(value)


def test_missing_directory_and_permissions(tmp_path, monkeypatch):
    with pytest.raises(ValueError, match="不存在"):
        read_text_file(str(tmp_path / "missing"))
    with pytest.raises(ValueError, match="普通文件"):
        read_text_file(str(tmp_path))
    path = tmp_path / "existing.txt"
    path.write_text("内容", encoding="utf-8")
    monkeypatch.setattr(type(path), "open", lambda *a, **kw: (_ for _ in ()).throw(PermissionError()))
    with pytest.raises(ValueError, match="权限"):
        read_text_file(str(path))


@pytest.mark.parametrize("model,args", [
    (TextFile, {"offset": 0}), (TextFile, {"limit": 1001}), (TextFile, {"limit": True}),
    (TextFile, {"offset": "1"}), (TextFile, {"encoding": ""}),
    (DocumentFile, {"offset": -1}), (DocumentFile, {"limit": 30001}),
    (DocumentFile, {"limit": 0}), (DocumentFile, {"unknown": True}),
])
def test_strict_arguments(model, args):
    with pytest.raises(ValidationError):
        model.model_validate({"path": "unused", **args})


@pytest.mark.parametrize("format", ["docx", "xlsx"])
async def test_real_office_worker_and_lossless_character_pagination(tmp_path, format):
    path = (make_docx if format == "docx" else make_xlsx)(tmp_path / f"资料.{format.upper()}")
    original = path.read_bytes()
    whole = await read_document(str(path))
    markdown = whole["markdown"]
    assert "林澈" in markdown and "文档完整尾部" in markdown
    if format == "docx":
        assert "# 世界设定" in markdown and "* 航海经历" in markdown and "|" in markdown
    else:
        assert markdown.index("## 人物") < markdown.index("## 空表") < markdown.index("## 世界")
        assert "有缓存公式" in markdown and "3" in markdown and "无缓存公式" in markdown
        assert "=1+2" not in markdown and "=2+3" not in markdown
    # Exercise the production worker once above; repeat pure conversion for small-page coverage.
    pages, offset = [], 0
    while True:
        result = convert_document(str(path), offset, 23)
        pages.append(result["markdown"])
        assert result["returned_chars"] == len(result["markdown"]) <= 23
        assert result["total_chars"] == len(markdown)
        if not result["has_more"]:
            assert result["next_offset"] is None
            break
        offset = result["next_offset"]
    assert "".join(pages) == markdown
    for offset in (len(markdown), len(markdown) + 10):
        empty = convert_document(str(path), offset, 23)
        assert empty["markdown"] == "" and not empty["has_more"] and empty["next_offset"] is None
    assert path.read_bytes() == original


def test_document_archive_and_output_limits(tmp_path, monkeypatch):
    path = make_docx(tmp_path / "资料.docx")
    blob = path.read_bytes()
    with monkeypatch.context() as patch:
        patch.setattr(editor_files, "MAX_ZIP_ENTRIES", 1)
        with pytest.raises(ValueError, match="4096"):
            convert_document(str(path), 0, 12_000)
    with monkeypatch.context() as patch:
        patch.setattr(editor_files, "MAX_EXPANDED_BYTES", 10)
        with pytest.raises(ValueError, match="128 MiB"):
            convert_document(str(path), 0, 12_000)
    with monkeypatch.context() as patch:
        patch.setattr(editor_files, "MAX_DOCUMENT_CHARS", 2)
        with pytest.raises(ValueError, match="800 万"):
            convert_document(str(path), 0, 12_000)
    # Mark a genuine ZIP member as encrypted in the central directory.
    data = bytearray(blob)
    pos = data.index(b"PK\x01\x02")
    data[pos + 8] |= 1
    path.write_bytes(data)
    with pytest.raises(ValueError, match="加密"):
        convert_document(str(path), 0, 12_000)
    path.write_bytes(b"not an OOXML document")
    with pytest.raises(ValueError, match="损坏"):
        convert_document(str(path), 0, 12_000)


async def test_worker_errors_are_readable(tmp_path):
    path = tmp_path / "broken.docx"
    path.write_bytes(b"encrypted or corrupt")
    with pytest.raises(ValueError, match="损坏|加密"):
        await read_document(str(path))
    other = tmp_path / "book.xls"
    other.write_bytes(b"legacy")
    with pytest.raises(ValueError, match="仅支持"):
        await read_document(str(other))
    with pytest.raises(ValueError, match="不存在"):
        await read_document(str(tmp_path / "missing.xlsx"))


@pytest.mark.parametrize("scenario", ["timeout", "cancel", "cancel_launch", "crash"])
async def test_worker_is_reaped_on_timeout_cancel_and_crash(monkeypatch, scenario):
    real_spawn = asyncio.create_subprocess_exec
    launched, release = asyncio.Event(), asyncio.Event()
    processes = []

    async def spawn(*args, **kwargs):
        program = "import sys; sys.exit(4)" if scenario == "crash" else "import time; time.sleep(60)"
        process = await real_spawn(sys.executable, "-c", program, **kwargs)
        processes.append(process)
        launched.set()
        if scenario == "cancel_launch":
            await release.wait()
        return process

    monkeypatch.setattr(asyncio, "create_subprocess_exec", spawn)
    if scenario == "timeout":
        monkeypatch.setattr(editor_files, "CONVERSION_TIMEOUT", 0.05)
    task = asyncio.create_task(read_document("unused"))
    await asyncio.wait_for(launched.wait(), 10)
    if scenario.startswith("cancel"):
        task.cancel()
        release.set()
        with pytest.raises(asyncio.CancelledError):
            await asyncio.wait_for(task, 10)
    else:
        with pytest.raises(ValueError, match="30 秒" if scenario == "timeout" else "异常退出"):
            await asyncio.wait_for(task, 10)
    assert processes[0].returncode is not None


async def test_short_path_handles_restore_without_granting_draft_reads(tmp_path):
    path = tmp_path / "12345678-1234-1234-1234-123456789abc.txt"
    path.write_text("世界观资料", encoding="utf-8")
    view = ModelView(str(path))
    owner = SimpleNamespace(context=SimpleNamespace(view=view), legacy=False, reads={})
    result = await EditorTools(owner).execute("read_text_file", {"path": view.text(str(path))})
    assert result["text"] == "世界观资料"
    assert "12345678-1234" not in result["path"]
    assert view.restore(result["path"]) == str(path)
    assert owner.reads == {}


@pytest.mark.parametrize("api", ["anthropic", "responses"])
async def test_file_pages_reach_model_intact_with_errors_and_native_schema(tmp_path, monkeypatch, api):
    path = tmp_path / "参考.txt"
    text = "长篇参考" * 1800 + "第一页尾部\n"
    path.write_bytes((text + "第二页").encode("utf-8"))
    document = make_docx(tmp_path / "资料.docx")
    expected_markdown = convert_document(str(document), 0, 30)["markdown"]
    editor = Editor(tmp_path / "store", monkeypatch, api)
    try:
        result = await editor.run([
            editor.tool(1, "read_text_file", {"path": str(path), "limit": 1}),
            editor.tool(2, "read_document", {"path": str(document), "limit": 30}),
            editor.tool(3, "read_document", {"path": str(tmp_path / "missing.docx")}),
            blocks(api, "读取完成")])
        body = json.loads(editor.requests[1].content)
        wire = editor.wire(editor.requests[1])
        results = ([p for m in wire for p in m.get("content", [])
                    if isinstance(p, dict) and p.get("type") == "tool_result"] if api == "anthropic"
                   else [m for m in wire if m.get("type") == "function_call_output"])
        page = json.loads(results[0]["content" if api == "anthropic" else "output"])
        assert page["text"] == text and page["has_more"] and page["next_offset"] == 2
        for name in ("read_text_file", "read_document"):
            schema = next(t for t in body["tools"] if t["name"] == name)
            parameters = schema["input_schema" if api == "anthropic" else "parameters"]
            assert parameters["additionalProperties"] is False
            assert "offset" in parameters["properties"] and "limit" in parameters["properties"]
        doc_wire = editor.wire(editor.requests[2])
        doc_results = ([p for m in doc_wire for p in m.get("content", [])
                        if isinstance(p, dict) and p.get("type") == "tool_result"] if api == "anthropic"
                       else [m for m in doc_wire if m.get("type") == "function_call_output"])
        doc_page = json.loads(doc_results[-1]["content" if api == "anthropic" else "output"])
        assert doc_page["markdown"] == expected_markdown
        assert doc_page["next_offset"] == 30 and doc_page["has_more"]
        last = editor.wire(editor.requests[3])
        if api == "anthropic":
            error_result = [p for m in last for p in m.get("content", [])
                            if isinstance(p, dict) and p.get("type") == "tool_result"][-1]
            assert error_result["is_error"] is True
        assert "文件不存在" in json.dumps(last, ensure_ascii=False)
        events = editor.service.edits.events(result["id"])
        activity = next(e["payload"] for e in events if e["kind"] == "tool_finished"
                        and e["payload"]["tool"] == "read_text_file")
        assert activity["result_truncated"] and "第一页尾部" not in activity["result"]
        assert editor.service.edits.context(editor.workspace["id"])["read_versions"] == {}
    finally:
        await editor.service.close()


async def test_resume_uses_saved_page_without_rereading_changed_file(tmp_path):
    path = tmp_path / "参考.txt"
    path.write_text("首次读取的内容", encoding="utf-8")
    store = Store(tmp_path / "data")
    editor = CardEditorService(store)
    workspace = editor.edits.create(import_fixture(store)["id"])

    class Failure(ToolLLM):
        async def chat(self, messages, **kwargs):
            if self.call_count == 1:
                raise ValueError("offline transport failure")
            async for chunk in super().chat(messages, **kwargs):
                yield chunk

    try:
        failed = await run(editor, workspace["id"], Failure([("read_text_file", {"path": str(path)})]))
        assert failed["status"] == "error"
        path.write_text("源文件已改变", encoding="utf-8")
        await editor.close()
        editor = CardEditorService(Store(tmp_path / "data"))
        provider = ToolLLM([])
        result = await run(editor, workspace["id"], provider, "", resume=failed["id"])
        assert result["status"] == "completed"
        messages = json.dumps(provider.call_log, ensure_ascii=False)
        assert "首次读取的内容" in messages and "源文件已改变" not in messages
        assert not any(e["kind"] == "tool_started" for e in editor.edits.events(result["id"]))
    finally:
        await editor.close()
